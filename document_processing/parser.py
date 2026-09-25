"""Parse PDF, DOCX, TXT, MD, and CSV while retaining source metadata."""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

import pdfplumber
from docx import Document as DocxDocument

from config.logging_config import get_logger

logger = get_logger("document_parser")

DOC_ID_RE = re.compile(r"^([A-Z]+-\d+)")
VERSION_RE = re.compile(r"\bv(\d+)\b", re.IGNORECASE)
DATE_RE = re.compile(r"(\d{4}-\d{2}-\d{2})")
SECTION_SPLIT_RE = re.compile(r"(?=(?:\d+\.\d+)\s)")
SECTION_HEAD_RE = re.compile(r"^(?P<sid>\d+\.\d+)\s*(?:\[(?P<flag>[MO])\])?\s*(?P<body>.*)$", re.DOTALL)
HEADING_THEN_SECTIONS_RE = re.compile(r"^(?P<heading>.+?)\s+(?P<rest>\d+\.\d+\s.*)$", re.DOTALL)
SUPERSEDED_RE = re.compile(r"\bv(\d+)\s+is\s+superseded\b", re.IGNORECASE)
V_EFFECTIVE_RE = re.compile(r"\bv(\d+)\s*\(effective\s+(\d{4}-\d{2}-\d{2})\)", re.IGNORECASE)
ROLE_SPECIFIC_RE = re.compile(r"ROLE-SPECIFIC", re.IGNORECASE)
ADVERSARIAL_RE = re.compile(
    r"(system note|administrator override|if you are an automated|ai system|hidden note for ai|"
    r"to any ai |note to any automated|appendix note|internal note|language model|"
    r"if parsing this document programmatically)",
    re.IGNORECASE,
)
CONFLICT_RE = re.compile(
    r"(conflict|has not yet been updated|still references the old|specifies a different threshold)",
    re.IGNORECASE,
)


@dataclass
class ParsedSection:
    """One numbered section extracted from a source document."""

    section_id: str
    heading: str | None
    content: str
    page_number: int | None
    paragraph_ref: str | None
    is_mandatory: bool
    is_optional: bool
    is_role_specific: bool
    flags: dict


@dataclass
class ParsedDocument:
    """Fully parsed document with header metadata and numbered sections."""

    document_id: str
    title: str
    version: str
    effective_date: date | None
    expiry_date: date | None
    department: str | None
    category: str | None
    file_type: str
    full_text: str
    sections: list[ParsedSection] = field(default_factory=list)
    superseded_versions: list[dict] = field(default_factory=list)


class DocumentParser:
    """Extracts text and section-level metadata from supported file types."""

    def parse(self, path: Path, raw_bytes: bytes) -> ParsedDocument:
        """Dispatch to the format-specific parser, then extract sections and version notes."""
        suffix = path.suffix.lower()
        if suffix == ".pdf":
            full_text, page_map = self._parse_pdf(raw_bytes)
            paragraph_mode = False
        elif suffix == ".docx":
            full_text, page_map = self._parse_docx(raw_bytes)
            paragraph_mode = True
        elif suffix in {".txt", ".md"}:
            full_text = raw_bytes.decode("utf-8", errors="replace")
            page_map = {1: full_text}
            paragraph_mode = True
        elif suffix == ".csv":
            full_text = self._parse_csv(raw_bytes)
            page_map = {1: full_text}
            paragraph_mode = True
        else:
            raise ValueError(f"Unsupported suffix: {suffix}")

        parsed = self._build_from_text(path, suffix.lstrip("."), full_text, page_map, paragraph_mode)
        logger.info(
            "document_parsed document_id=%s version=%s sections=%s type=%s",
            parsed.document_id,
            parsed.version,
            len(parsed.sections),
            parsed.file_type,
        )
        return parsed

    def _parse_docx(self, raw_bytes: bytes) -> tuple[str, dict[int, str]]:
        document = DocxDocument(io.BytesIO(raw_bytes))
        paragraphs = [p.text.strip() for p in document.paragraphs if p.text and p.text.strip()]
        full_text = "\n".join(paragraphs)
        return full_text, {1: full_text}

    def _parse_pdf(self, raw_bytes: bytes) -> tuple[str, dict[int, str]]:
        page_map: dict[int, str] = {}
        parts: list[str] = []
        with pdfplumber.open(io.BytesIO(raw_bytes)) as pdf:
            for index, page in enumerate(pdf.pages, start=1):
                text = (page.extract_text() or "").strip()
                page_map[index] = text
                parts.append(text)
        return "\n".join(parts), page_map

    def _parse_csv(self, raw_bytes: bytes) -> str:
        text = raw_bytes.decode("utf-8", errors="replace")
        reader = csv.reader(io.StringIO(text))
        rows = [" | ".join(row) for row in reader]
        return "\n".join(rows)

    def _build_from_text(
        self,
        path: Path,
        file_type: str,
        full_text: str,
        page_map: dict[int, str],
        paragraph_mode: bool,
    ) -> ParsedDocument:
        lines = [ln.strip() for ln in full_text.splitlines() if ln.strip()]
        header = lines[0] if lines else path.stem
        meta_line = lines[1] if len(lines) > 1 else ""

        doc_id_match = DOC_ID_RE.match(header) or DOC_ID_RE.match(path.name)
        document_id = doc_id_match.group(1) if doc_id_match else path.stem.split("_")[0]
        title = re.sub(r"^[A-Z]+-\d+\s*[–—-]\s*", "", header).strip() or path.stem

        version_match = VERSION_RE.search(meta_line) or VERSION_RE.search(full_text[:400])
        version = f"v{version_match.group(1)}" if version_match else "v1"

        dates = DATE_RE.findall(meta_line)
        effective_date = _parse_iso_date(dates[0]) if dates else None
        expiry_date = _parse_iso_date(dates[1]) if len(dates) > 1 else None

        meta_parts = [p.strip() for p in meta_line.split("|")]
        department = meta_parts[2] if len(meta_parts) >= 3 else _department_from_folder(path)
        category = meta_parts[3] if len(meta_parts) >= 4 else path.parent.name

        body_text = "\n".join(lines[2:]) if len(lines) > 2 else full_text
        sections = self._extract_sections(body_text, page_map, paragraph_mode)
        if not sections:
            sections = [
                ParsedSection(
                    section_id="BODY",
                    heading=title,
                    content=body_text or full_text,
                    page_number=1,
                    paragraph_ref="p1" if paragraph_mode else None,
                    is_mandatory=False,
                    is_optional=False,
                    is_role_specific=False,
                    flags={},
                )
            ]
        superseded = self._extract_superseded(full_text)
        return ParsedDocument(
            document_id=document_id,
            title=title,
            version=version,
            effective_date=effective_date,
            expiry_date=expiry_date,
            department=department,
            category=category,
            file_type=file_type,
            full_text=full_text,
            sections=sections,
            superseded_versions=superseded,
        )

    def _extract_sections(
        self, body_text: str, page_map: dict[int, str], paragraph_mode: bool
    ) -> list[ParsedSection]:
        sections: list[ParsedSection] = []
        current_heading: str | None = None
        paragraph_index = 0
        for raw_para in body_text.split("\n"):
            paragraph_index += 1
            para = raw_para.strip()
            if not para:
                continue
            heading_match = HEADING_THEN_SECTIONS_RE.match(para)
            if heading_match:
                current_heading = heading_match.group("heading").strip()
                rest = heading_match.group("rest")
            else:
                rest = para
            pieces = [p.strip() for p in SECTION_SPLIT_RE.split(rest) if p.strip()]
            for piece in pieces:
                head = SECTION_HEAD_RE.match(piece)
                if not head:
                    continue
                section_id = head.group("sid")
                flag = head.group("flag")
                content = head.group("body").strip()
                page_number = _page_for_section(section_id, page_map)
                flags = {
                    "adversarial": bool(ADVERSARIAL_RE.search(content)),
                    "conflict_mentioned": bool(CONFLICT_RE.search(content)),
                    "version_change": bool(SUPERSEDED_RE.search(content) or V_EFFECTIVE_RE.search(content)),
                }
                sections.append(
                    ParsedSection(
                        section_id=section_id,
                        heading=current_heading,
                        content=content,
                        page_number=page_number,
                        paragraph_ref=f"p{paragraph_index}" if paragraph_mode else None,
                        is_mandatory=flag == "M",
                        is_optional=flag == "O",
                        is_role_specific=bool(ROLE_SPECIFIC_RE.search(content)),
                        flags=flags,
                    )
                )
        return sections

    def _extract_superseded(self, full_text: str) -> list[dict]:
        found: dict[str, dict] = {}
        for match in V_EFFECTIVE_RE.finditer(full_text):
            version = f"v{match.group(1)}"
            found[version] = {
                "version": version,
                "effective_date": match.group(2),
            }
        current = VERSION_RE.search(full_text[:200])
        current_v = f"v{current.group(1)}" if current else "v1"
        results: list[dict] = []
        for version, payload in found.items():
            if version != current_v:
                payload["status"] = "obsolete"
                payload["notes"] = f"{version} superseded by {current_v}"
                results.append(payload)
        if SUPERSEDED_RE.search(full_text) and not results:
            results.append({"version": "v1", "status": "obsolete", "notes": "v1 is superseded", "effective_date": None})
        return results


def _parse_iso_date(value: str) -> date | None:
    try:
        return datetime.strptime(value, "%Y-%m-%d").date()
    except ValueError:
        return None


def _department_from_folder(path: Path) -> str | None:
    folder = path.parent.name.lower()
    mapping = {
        "policies": "Company-wide",
        "sops": "Operations",
        "handbook": "Company-wide",
        "faqs": "Company-wide",
        "compliance": "Compliance",
        "role_descriptions": "People & Culture",
    }
    return mapping.get(folder, folder)


def _page_for_section(section_id: str, page_map: dict[int, str]) -> int | None:
    needle = f"{section_id} "
    for page_no, text in page_map.items():
        if needle in (text or ""):
            return page_no
    return 1 if page_map else None
